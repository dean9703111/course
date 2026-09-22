#!/usr/bin/env python3
"""讀學員檔「完工資料」，依故障代碼畫兩張柏拉圖（維修工時／工單次數）到對照答案。

跑法：
    uv run --with openpyxl --with matplotlib --with pillow python3 _scripts/pareto_erp.py
"""

import shutil
from collections import defaultdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl

plt.rcParams["font.family"] = "Heiti TC"

TAHOHO_DIR = Path(__file__).resolve().parent.parent
STUDENT_XLSX = TAHOHO_DIR / "example" / "cowork-erp-kpi" / "2024_B廠_ERP維護績效.xlsx"
ANSWER_DIR = TAHOHO_DIR / "example" / "cowork-erp-kpi" / "對照答案"
ASSETS_DIR = TAHOHO_DIR / "assets"

BAR_HIGHLIGHT = "#d9480f"
BAR_DIM = "#adb5bd"
LINE_COLOR = "#1c7ed6"


def wrap_label(label, width=6):
    if len(label) <= width:
        return label
    return label[:width] + "\n" + label[width:]


def load_groups():
    wb = openpyxl.load_workbook(STUDENT_XLSX, data_only=True)
    ws = wb["完工資料"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_code = header.index("故障代碼")
    idx_repair_min = header.index("維修時間(分)")

    count_by = defaultdict(int)
    hours_by = defaultdict(float)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[0].value is None:
            continue
        code = row[idx_code].value or "未填"
        count_by[code] += 1
        hours_by[code] += (row[idx_repair_min].value or 0) / 60
    return count_by, hours_by


def plot_pareto(data, title, ylabel, out_path):
    items = sorted(data.items(), key=lambda kv: kv[1], reverse=True)
    labels = [k for k, _v in items]
    values = [v for _k, v in items]
    total = sum(values)

    cum = []
    running = 0.0
    for v in values:
        running += v
        cum.append(running / total * 100)

    n80 = next((i + 1 for i, c in enumerate(cum) if c >= 80), len(cum))
    pct80 = cum[n80 - 1]

    colors = [BAR_HIGHLIGHT if i < n80 else BAR_DIM for i in range(len(cum))]
    wrapped_labels = [wrap_label(l) for l in labels]

    fig, ax1 = plt.subplots(figsize=(11, 6), dpi=130)
    bars = ax1.bar(wrapped_labels, values, color=colors)
    ax1.set_ylabel(ylabel)
    ax1.set_title(title)

    for bar, v in zip(bars, values):
        ax1.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height(),
            f"{v:.0f}",
            ha="center",
            va="bottom",
            fontsize=9,
        )

    plt.setp(ax1.get_xticklabels(), rotation=30, ha="right", fontsize=9)

    ax2 = ax1.twinx()
    ax2.plot(wrapped_labels, cum, color=LINE_COLOR, marker="o")
    ax2.axhline(80, color=LINE_COLOR, linestyle="--", linewidth=1)
    ax2.set_ylim(0, 105)
    ax2.set_ylabel("累積百分比 (%)")

    ax1.annotate(
        f"前 {n80} 項原因佔 {pct80:.0f}%",
        xy=(0.98, 0.02),
        xycoords="axes fraction",
        ha="right",
        va="bottom",
        fontsize=12,
        color=BAR_HIGHLIGHT,
    )

    fig.subplots_adjust(bottom=0.30)
    fig.savefig(out_path)
    plt.close(fig)

    print(f"\n{title}")
    print("前 5 名與累積百分比：")
    for i in range(min(5, len(items))):
        print(f"  {labels[i]}: {values[i]:.1f}  累積 {cum[i]:.1f}%")
    print(f"前 {n80} 項累積達 {pct80:.1f}%（首次 >= 80%）")


def main():
    count_by, hours_by = load_groups()

    ANSWER_DIR.mkdir(parents=True, exist_ok=True)

    plot_pareto(
        hours_by,
        "B廠 2024 年 1 至 4 月故障原因柏拉圖（維修工時）",
        "維修工時（小時）",
        ANSWER_DIR / "柏拉圖_維修工時.png",
    )
    plot_pareto(
        count_by,
        "B廠 2024 年 1 至 4 月故障原因柏拉圖（工單次數）",
        "工單次數",
        ANSWER_DIR / "柏拉圖_工單次數.png",
    )

    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    src = ANSWER_DIR / "柏拉圖_維修工時.png"
    dst = ASSETS_DIR / "erp-pareto-hours.png"
    shutil.copy(src, dst)
    print(f"\n已複製 {src} -> {dst}")


if __name__ == "__main__":
    main()
