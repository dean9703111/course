#!/usr/bin/env python3
"""讀學員檔（2024_B廠_ERP維護績效.xlsx），合成 2024_C廠_ERP維護績效.xlsx 供多廠匯入練習。

random seed 固定 42，可重跑得到一致結果。

跑法：
    uv run --with openpyxl --with matplotlib --with pillow python3 _scripts/synth_plant_c.py
"""

import csv
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from deidentify_erp import (
    load_extra_terms,
    verify_abspath_and_refreshed_by,
    verify_no_pivot_entries,
    verify_no_taho,
    verify_terms_zero,
)

TAHOHO_DIR = Path(__file__).resolve().parent.parent
EXAMPLE_DIR = TAHOHO_DIR / "example" / "cowork-erp-kpi"
STUDENT_XLSX = EXAMPLE_DIR / "2024_B廠_ERP維護績效.xlsx"
PLANT_C_XLSX = EXAMPLE_DIR / "2024_C廠_ERP維護績效.xlsx"
NAME_MAP_CSV = TAHOHO_DIR / "_scripts" / "erp_name_map.csv"

SEED = 42

# 各欄位「前綴」替換規則：只動符合前綴的部分，其餘字串不動
PREFIX_MAP = {
    "派工單號": {"MBA": "MCA", "MBL": "MCL"},
    "請修單號": {"LBA": "LCA", "LBL": "LCL"},
    "設備編號": {"BA0": "CA0"},
}
CLEAR_COLS = ["MTTF", "稼動率%", "妥善率%"]


def replace_prefix(value, prefix_map):
    for old, new in prefix_map.items():
        if value.startswith(old):
            return new + value[len(old):]
    return value


def transform_value(v, col_name):
    if isinstance(v, str):
        if col_name in PREFIX_MAP:
            v = replace_prefix(v, PREFIX_MAP[col_name])
        if "B廠" in v:
            v = v.replace("B廠", "C廠")
    return v


def load_sheet_rows(ws):
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [list(r) for r in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True)]
    return header, rows


def build_wangong(ws_src, rng):
    """完工資料：隨機丟 15% 列，維修/停機時間依同一係數縮放，MTTF 等三欄清空。"""
    header, rows = load_sheet_rows(ws_src)
    idx_repair = header.index("維修時間(分)")
    idx_down = header.index("停機時間(分)")
    clear_idx = [header.index(c) for c in CLEAR_COLS if c in header]

    n_drop = round(len(rows) * 0.15)
    drop_idx = set(rng.sample(range(len(rows)), n_drop))

    out_rows = []
    for i, row in enumerate(rows):
        if i in drop_idx:
            continue
        row = [transform_value(v, header[j]) for j, v in enumerate(row)]
        factor = rng.uniform(0.7, 1.3)
        for idx in (idx_repair, idx_down):
            v = row[idx]
            if isinstance(v, (int, float)):
                row[idx] = max(5, round(v * factor / 5) * 5)
        for idx in clear_idx:
            if row[idx] is not None:
                row[idx] = None
        out_rows.append(row)
    return header, out_rows


def build_weiwangong(ws_src):
    """未完工資料：不丟列，只做字串替換。"""
    header, rows = load_sheet_rows(ws_src)
    out_rows = []
    for row in rows:
        out_rows.append([transform_value(v, header[j]) for j, v in enumerate(row)])
    return header, out_rows


def build_data(ws_src, rng):
    """DATA：先抽 15% 的派工單號，同單號的列一起丟。"""
    header, rows = load_sheet_rows(ws_src)
    idx_ticket = header.index("派工單號")

    unique_tickets = sorted({row[idx_ticket] for row in rows if row[idx_ticket] is not None})
    n_drop = round(len(unique_tickets) * 0.15)
    drop_tickets = set(rng.sample(unique_tickets, n_drop))

    out_rows = []
    for row in rows:
        if row[idx_ticket] in drop_tickets:
            continue
        out_rows.append([transform_value(v, header[j]) for j, v in enumerate(row)])
    return header, out_rows, len(unique_tickets), len(drop_tickets)


def write_sheet(wb_out, sheet_name, header, rows, ws_src):
    ws_out = wb_out.create_sheet(sheet_name)
    ws_out.append(header)
    for row in rows:
        ws_out.append(row)
    bold = Font(bold=True)
    for cell in ws_out[1]:
        cell.font = bold
    for col_letter, dim in ws_src.column_dimensions.items():
        if dim.width:
            ws_out.column_dimensions[col_letter].width = dim.width


def main():
    wb_src = openpyxl.load_workbook(STUDENT_XLSX, data_only=True)
    rng = random.Random(SEED)

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    header_w, rows_w = build_wangong(wb_src["完工資料"], rng)
    write_sheet(wb_out, "完工資料", header_w, rows_w, wb_src["完工資料"])

    header_u, rows_u = build_weiwangong(wb_src["未完工資料"])
    write_sheet(wb_out, "未完工資料", header_u, rows_u, wb_src["未完工資料"])

    header_d, rows_d, n_unique, n_drop = build_data(wb_src["DATA"], rng)
    write_sheet(wb_out, "DATA", header_d, rows_d, wb_src["DATA"])

    wb_out.save(PLANT_C_XLSX)

    names = []
    with open(NAME_MAP_CSV, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            names.append(row["真名"])
    extra_terms = [old for old, _new in load_extra_terms()]
    all_terms = names + extra_terms

    idx_repair_w = header_w.index("維修時間(分)")
    idx_down_w = header_w.index("停機時間(分)")
    total_repair_hours = sum(r[idx_repair_w] for r in rows_w if r[idx_repair_w] is not None) / 60
    total_down_hours = sum(r[idx_down_w] for r in rows_w if r[idx_down_w] is not None) / 60

    print(f"已寫入 {PLANT_C_XLSX}")
    print(f"C廠 完工資料列數：{len(rows_w)}（原始 334，丟棄 {334 - len(rows_w)} 列）")
    print(f"C廠 完工資料 維修工時合計：{total_repair_hours:.2f} 小時")
    print(f"C廠 完工資料 停機合計：{total_down_hours:.2f} 小時")
    print(f"C廠 DATA 原始單號數 {n_unique}，丟棄 {n_drop} 個單號，剩餘列數 {len(rows_d)}")

    print("\n===== 驗證 =====")
    hits = verify_terms_zero(PLANT_C_XLSX, all_terms)
    if hits:
        raise RuntimeError(f"C廠檔仍殘留真名／詞表字串：{hits}")
    print(f"真名＋詞表殘留次數：0（檢查 {len(all_terms)} 個字串）")

    abspath_count, refreshed_by_values = verify_abspath_and_refreshed_by(PLANT_C_XLSX)
    print(f"x15ac:absPath 出現次數：{abspath_count}")
    print(f"refreshedBy 值：{refreshed_by_values}")
    if abspath_count != 0:
        raise RuntimeError(f"C廠檔仍殘留 x15ac:absPath：{abspath_count} 次")
    if refreshed_by_values:
        raise RuntimeError(f"C廠檔不應含 refreshedBy 屬性：{refreshed_by_values}")

    taho_count = verify_no_taho(PLANT_C_XLSX)
    print(f"taho 子字串（不分大小寫）命中次數：{taho_count}")
    if taho_count != 0:
        raise RuntimeError(f"C廠檔仍殘留 taho 子字串：{taho_count} 次")

    pivot_entries = verify_no_pivot_entries(PLANT_C_XLSX)
    print(f"xl/pivotCache 或 xl/pivotTables 項目數：{len(pivot_entries)}")
    if pivot_entries:
        raise RuntimeError(f"C廠檔不應含樞紐表項目：{pivot_entries}")

    if len(rows_w) != 284:
        raise RuntimeError(f"C廠完工資料列數不符：{len(rows_w)} != 284")
    print(f"C廠 完工資料列數確認：{len(rows_w)}（預期 284）")

    print("\n全部驗證通過。")


if __name__ == "__main__":
    main()
