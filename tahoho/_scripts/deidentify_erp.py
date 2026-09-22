#!/usr/bin/env python3
"""去識別化 2024_BL_ERP維護績效.xlsx，產出課程範例包。

跑法：
    uv run --with openpyxl --with matplotlib --with pillow python3 _scripts/deidentify_erp.py

輸出：
  - _scripts/erp_name_map.csv                                      真名／代號對照表（不進範例包）
  - example/cowork-erp-kpi/2024_B廠_ERP維護績效.xlsx                學員檔（完工資料/未完工資料/DATA）
  - example/cowork-erp-kpi/對照答案/人工樞紐統計.xlsx                其餘 21 張樞紐表（靜態值）
  - example/cowork-needs-folder/2024_BL_ERP維護績效.xlsx（覆寫）     保留樞紐表的課前需求包同名檔
"""

import csv
import html
import re
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

SRC = Path("/Users/lindingyuan/Downloads/cowork-needs-folder/2024_BL_ERP維護績效.xlsx")
TAHOHO_DIR = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = TAHOHO_DIR / "_scripts"
EXAMPLE_DIR = TAHOHO_DIR / "example" / "cowork-erp-kpi"
ANSWER_DIR = EXAMPLE_DIR / "對照答案"
NEEDS_FOLDER_XLSX = TAHOHO_DIR / "example" / "cowork-needs-folder" / "2024_BL_ERP維護績效.xlsx"

STUDENT_XLSX = EXAMPLE_DIR / "2024_B廠_ERP維護績效.xlsx"
ANSWER_XLSX = ANSWER_DIR / "人工樞紐統計.xlsx"
NAME_MAP_CSV = SCRIPTS_DIR / "erp_name_map.csv"
EXTRA_TERMS_CSV = SCRIPTS_DIR / "erp_extra_terms.csv"

DETAIL_SHEETS = ["完工資料", "未完工資料", "DATA"]
PLACE_OLD, PLACE_NEW = "八里", "B廠"


# ---------- 名單收集與代號編號 ----------

def collect_name_map(wb):
    """從 DATA 分頁的「維修人員」欄收集真名，依組別＋首次出現順序編號成代號。"""
    ws = wb["DATA"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_name = header.index("維修人員")
    idx_group = header.index("組別")

    order = []
    seen = set()
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        name = row[idx_name].value
        if not name or name in seen:
            continue
        seen.add(name)
        order.append((name, row[idx_group].value))

    counters = {}
    mapping = {}
    csv_rows = []
    for name, group in order:
        n = counters.get(group, 0) + 1
        counters[group] = n
        code = f"{group}-{n:02d}"
        mapping[name] = code
        csv_rows.append((name, code, group))
    return mapping, csv_rows


def write_name_map_csv(csv_rows):
    SCRIPTS_DIR.mkdir(parents=True, exist_ok=True)
    with open(NAME_MAP_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["真名", "代號", "組別"])
        w.writerows(csv_rows)


def load_extra_terms():
    """讀額外替換詞表（人名/暱稱/車牌/地點/機構/品牌/租戶等），不進範例包。"""
    terms = []
    with open(EXTRA_TERMS_CSV, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            terms.append((row["原字串"], row["替換"]))
    return terms


def build_pattern(mapping, extra_terms, include_place):
    """長字串優先排序，回傳 (regex, lookup dict)，單一 pass 避免替換結果被二次改寫。

    不分大小寫（含 taho 系列租戶字串與英文人名的大小寫變體）。
    """
    pairs = list(mapping.items()) + list(extra_terms)
    if include_place:
        pairs.append((PLACE_OLD, PLACE_NEW))
    pairs.sort(key=lambda kv: len(kv[0]), reverse=True)
    pattern = re.compile("|".join(re.escape(old) for old, _new in pairs), re.IGNORECASE)
    lookup = {old.lower(): new for old, new in pairs}
    return pattern, lookup


def replace_str(value, pattern, lookup):
    return pattern.sub(lambda m: lookup[m.group(0).lower()], value)


# ---------- 學員檔（完工資料／未完工資料／DATA） ----------

def detect_last_row(ws):
    """以 A 欄是否有值判斷實際資料列數的最後一列（去掉尾端空白格式列）。"""
    last_row = 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[0].value is not None:
            last_row = row[0].row
    return last_row


def build_student_workbook(wb_src, pattern, lookup):
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    wb_out.properties.creator = "corp"
    bold = Font(bold=True)

    for sheet_name in DETAIL_SHEETS:
        ws_src = wb_src[sheet_name]
        last_row = detect_last_row(ws_src)
        ws_out = wb_out.create_sheet(sheet_name)

        for row in ws_src.iter_rows(min_row=1, max_row=last_row):
            values = []
            for cell in row:
                v = cell.value
                if isinstance(v, str):
                    v = replace_str(v, pattern, lookup)
                values.append(v)
            ws_out.append(values)

        for cell in ws_out[1]:
            cell.font = bold

        for col_letter, dim in ws_src.column_dimensions.items():
            if dim.width:
                ws_out.column_dimensions[col_letter].width = dim.width

    EXAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    wb_out.save(STUDENT_XLSX)


# ---------- 對照答案（其餘 21 張樞紐表） ----------

def build_answer_workbook(wb_src, pattern, lookup):
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    wb_out.properties.creator = "corp"

    for sheet_name in wb_src.sheetnames:
        if sheet_name in DETAIL_SHEETS:
            continue
        ws_src = wb_src[sheet_name]
        ws_out = wb_out.create_sheet(sheet_name)
        for row in ws_src.iter_rows(min_row=1, max_row=ws_src.max_row):
            values = []
            for cell in row:
                v = cell.value
                if isinstance(v, str):
                    v = replace_str(v, pattern, lookup)
                values.append(v)
            ws_out.append(values)

    ANSWER_DIR.mkdir(parents=True, exist_ok=True)
    wb_out.save(ANSWER_XLSX)


# ---------- 課前需求包同名檔（zip 層級，保留樞紐表） ----------

CORE_XML_NAME = "docProps/core.xml"
WORKBOOK_XML_NAME = "xl/workbook.xml"


def scrub_core_xml(text):
    """把 docProps/core.xml 的 dc:creator、cp:lastModifiedBy 內容改成 corp（保留標籤）。"""
    text = re.sub(r"(<dc:creator[^>]*>)[^<]*(</dc:creator>)", r"\g<1>corp\g<2>", text)
    text = re.sub(
        r"(<cp:lastModifiedBy[^>]*>)[^<]*(</cp:lastModifiedBy>)", r"\g<1>corp\g<2>", text
    )
    return text


def scrub_workbook_xml(text):
    """刪掉 xl/workbook.xml 裡記錄本機路徑的 <x15ac:absPath .../> 元素。"""
    return re.sub(r"<x15ac:absPath[^>]*/>", "", text)


def scrub_pivot_cache_definition(text):
    """xl/pivotCache/pivotCacheDefinition*.xml 的 refreshedBy 一律改成 corp。"""
    return re.sub(r'refreshedBy="[^"]*"', 'refreshedBy="corp"', text)


def build_needs_folder_copy(name_pattern, name_lookup):
    NEEDS_FOLDER_XLSX.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(SRC, "r") as zin:
        infos = zin.infolist()
        with zipfile.ZipFile(NEEDS_FOLDER_XLSX, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in infos:
                data = zin.read(info.filename)
                if info.filename.endswith(".xml"):
                    # customXml/item1.xml 是 UTF-16（帶 BOM），其餘皆為 UTF-8
                    encoding = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8"
                    text = data.decode(encoding)
                    text = name_pattern.sub(lambda m: name_lookup[m.group(0).lower()], text)
                    if info.filename == CORE_XML_NAME:
                        text = scrub_core_xml(text)
                    elif info.filename == WORKBOOK_XML_NAME:
                        text = scrub_workbook_xml(text)
                    elif info.filename.startswith("xl/pivotCache/pivotCacheDefinition"):
                        text = scrub_pivot_cache_definition(text)
                    data = text.encode(encoding)
                info.compress_type = zipfile.ZIP_DEFLATED
                zout.writestr(info, data)


# ---------- 驗證 ----------

def zip_full_text(path):
    """把 zip 內所有 .xml 解碼串接成一份文字，供驗證用全文掃描。
    openpyxl 寫出的檔案會把中文存成 &#NNNNN; 數字實體，所以每個 part 都先 html.unescape，
    否則掃描會對中文字串視而不見。"""
    parts = []
    with zipfile.ZipFile(path) as z:
        for entry in z.namelist():
            if not entry.endswith(".xml"):
                continue
            data = z.read(entry)
            encoding = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8"
            parts.append(html.unescape(data.decode(encoding, errors="ignore")))
    return "\n".join(parts)


def verify_terms_zero(path, terms):
    """檢查一批字串（不分大小寫）在 zip 全文中的殘留次數，回傳非 0 的項目。"""
    text = zip_full_text(path).lower()
    hits = {}
    for term in terms:
        c = text.count(term.lower())
        if c:
            hits[term] = c
    return hits


def verify_abspath_and_refreshed_by(path):
    text = zip_full_text(path)
    abspath_count = len(re.findall(r"<x15ac:absPath[^>]*/>", text))
    refreshed_by_values = re.findall(r'refreshedBy="([^"]*)"', text)
    return abspath_count, refreshed_by_values


def verify_no_taho(path):
    """taho 子字串命中數（不分大小寫），排除 Excel 預設佈景主題內建字型名稱 Tahoma
    （原檔 xl/theme/theme1.xml 本來就有，與 tahotw/tahoho 租戶字串無關）。"""
    text = zip_full_text(path)
    return len(re.findall(r"taho(?!ma)", text, re.IGNORECASE))


def verify_no_pivot_entries(path):
    with zipfile.ZipFile(path) as z:
        return [
            n
            for n in z.namelist()
            if n.startswith("xl/pivotCache/") or n.startswith("xl/pivotTables/")
        ]


def verify_workbook_opens(path):
    wb = openpyxl.load_workbook(path)
    return [(name, wb[name].max_row) for name in wb.sheetnames]


def read_core_props(path):
    """讀 docProps/core.xml 的 creator／lastModifiedBy，供驗證印出改後的值。"""
    with zipfile.ZipFile(path) as z:
        if CORE_XML_NAME not in z.namelist():
            return {}
        data = z.read(CORE_XML_NAME)
        encoding = "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8"
        text = data.decode(encoding)
    props = {}
    m = re.search(r"<dc:creator[^>]*>([^<]*)</dc:creator>", text)
    if m:
        props["creator"] = m.group(1)
    m = re.search(r"<cp:lastModifiedBy[^>]*>([^<]*)</cp:lastModifiedBy>", text)
    if m:
        props["lastModifiedBy"] = m.group(1)
    return props


def main():
    print(f"讀取原檔：{SRC}")
    wb_src = openpyxl.load_workbook(SRC, data_only=True)

    mapping, csv_rows = collect_name_map(wb_src)
    write_name_map_csv(csv_rows)
    print(f"收集到 {len(mapping)} 個維修人員真名字串，已寫入 {NAME_MAP_CSV}")
    for name, code in mapping.items():
        print(f"  {name} -> {code}")

    extra_terms = load_extra_terms()
    print(f"\n讀取額外替換詞表 {EXTRA_TERMS_CSV}，共 {len(extra_terms)} 條")
    src_text_lower = zip_full_text(SRC).lower()
    print("原檔命中次數：")
    for old, new in extra_terms:
        c = src_text_lower.count(old.lower())
        print(f"  {old} -> {new}: {c}")

    all_extra_terms = [old for old, _new in extra_terms]

    pattern_with_place, lookup_with_place = build_pattern(
        mapping, extra_terms, include_place=True
    )
    pattern_name_only, lookup_name_only = build_pattern(
        mapping, extra_terms, include_place=False
    )

    print("\n建立學員檔...")
    build_student_workbook(wb_src, pattern_with_place, lookup_with_place)
    print(f"  已寫入 {STUDENT_XLSX}")

    print("\n建立對照答案（人工樞紐統計）...")
    build_answer_workbook(wb_src, pattern_with_place, lookup_with_place)
    print(f"  已寫入 {ANSWER_XLSX}")

    print("\n覆寫課前需求包同名檔（保留樞紐表）...")
    build_needs_folder_copy(pattern_name_only, lookup_name_only)
    print(f"  已寫入 {NEEDS_FOLDER_XLSX}")

    real_names = list(mapping.keys())
    all_terms = real_names + all_extra_terms

    print("\n===== 驗證 =====")
    for label, path, has_pivot in [
        ("學員檔", STUDENT_XLSX, False),
        ("對照答案", ANSWER_XLSX, False),
        ("課前需求包同名檔", NEEDS_FOLDER_XLSX, True),
    ]:
        print(f"\n-- {label}: {path}")

        hits = verify_terms_zero(path, all_terms)
        if hits:
            raise RuntimeError(f"{label} 仍殘留真名／詞表字串：{hits}")
        print(f"  真名＋詞表殘留次數：0（檢查 {len(all_terms)} 個字串）")

        sheets_info = verify_workbook_opens(path)
        print("  openpyxl 開啟成功，工作表與列數：")
        for name, max_row in sheets_info:
            print(f"    {name}: max_row={max_row}")

        core_props = read_core_props(path)
        print(f"  docProps/core.xml：{core_props}")

        abspath_count, refreshed_by_values = verify_abspath_and_refreshed_by(path)
        print(f"  x15ac:absPath 出現次數：{abspath_count}")
        print(f"  refreshedBy 值：{refreshed_by_values}")
        if abspath_count != 0:
            raise RuntimeError(f"{label} 仍殘留 x15ac:absPath：{abspath_count} 次")
        if has_pivot:
            if any(v != "corp" for v in refreshed_by_values):
                raise RuntimeError(f"{label} refreshedBy 未全改成 corp：{refreshed_by_values}")
        else:
            if refreshed_by_values:
                raise RuntimeError(f"{label} 不應含 refreshedBy 屬性：{refreshed_by_values}")

        taho_count = verify_no_taho(path)
        print(f"  taho 子字串（不分大小寫）命中次數：{taho_count}")
        if taho_count != 0:
            raise RuntimeError(f"{label} 仍殘留 taho 子字串：{taho_count} 次")

        if not has_pivot:
            pivot_entries = verify_no_pivot_entries(path)
            print(f"  xl/pivotCache 或 xl/pivotTables 項目數：{len(pivot_entries)}")
            if pivot_entries:
                raise RuntimeError(f"{label} 不應含樞紐表項目：{pivot_entries}")

    student_sheets = dict(verify_workbook_opens(STUDENT_XLSX))
    data_rows = student_sheets["DATA"] - 1
    wangong_rows = student_sheets["完工資料"] - 1
    print(f"\n學員檔 DATA 資料列數：{data_rows}（預期 14389）")
    print(f"學員檔 完工資料 資料列數：{wangong_rows}（預期 334）")
    if data_rows != 14389:
        raise RuntimeError(f"DATA 資料列數不符：{data_rows} != 14389")
    if wangong_rows != 334:
        raise RuntimeError(f"完工資料 資料列數不符：{wangong_rows} != 334")

    print("\n全部驗證通過。")


if __name__ == "__main__":
    main()
