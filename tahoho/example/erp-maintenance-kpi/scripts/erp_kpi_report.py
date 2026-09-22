#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
erp_kpi_report.py — ERP 維護績效月報產生器（以程式為準）
================================================================================
輸入：一份 .xlsx，只需三張工作表：
        DATA        ERP 派工單匯出（可含多年、可含重複列）
        完工資料     報表期間內已完工的派工單 ＋ 人工加註欄位
                     （設備編號、故障原因、故障代碼、完工報告、維修類別、
                       停機時間(分)、維修時間(分)、Work Plan）
        未完工資料   未結案派工單快照

輸出：一份 .xlsx，工作表名稱、順序、版面、圖表位置與「2024_BL_ERP維護績效.xlsx」
      範本相同（21 張分析表 ＋ 14 張圖 ＋ 3 張來源表），最後多一張
      「_資料校正紀錄」，列出程式覆寫過的每一個人工數值與原因。

以程式為準的意思
----------------
下列欄位一律由程式重算，輸入檔裡若已填值只拿來對帳、不採用：
    DATA.結單日數、DATA.符合度
    完工資料.MTTF、完工資料.稼動率%、完工資料.妥善率%
    （計算方式見 kpi_per_equipment()；公式已用 2024 年 1～4 月 185 台設備逐台驗證）
DATA 的完全重複列會先移除；完工資料重複的派工單號只留第一筆。

用法
----
    python erp_kpi_report.py 來源.xlsx
    python erp_kpi_report.py 來源.xlsx -o 輸出.xlsx
    python erp_kpi_report.py 來源.xlsx --period 2024-01-01 2024-04-30 --as-of 2024-05-24

需要：pandas、openpyxl（pip install pandas openpyxl）
"""

from __future__ import annotations

import argparse
import calendar
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, BarChart3D, PieChart, PieChart3D, RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ══════════════════════════════════════════════════════════════════════════════
#  CONFIG — 所有門檻與口徑集中在此
# ══════════════════════════════════════════════════════════════════════════════


@dataclass
class Config:
    # ── 報表期間（None ＝ 自動：完工資料「完成日」最多的那一年，1/1 起到最後一個月月底）
    period_start: date | None = None
    period_end: date | None = None
    # 未完工單「結單日數」算到哪一天（None ＝ 今天；要重現舊報表請指定）
    as_of: date | None = None

    # ── 維修類別。KPI 類報表只看 BD（故障維修）
    kpi_category: str = "BD"
    category_order: tuple[str, ...] = ("BD", "IM", "MR", "SI", "OH")
    # PLAN WORK 篩選：原檔只勾 IM、MR
    planwork_categories: tuple[str, ...] = ("IM", "MR")
    planwork_top_n: int = 10            # 前 N 名（含同分，與 Excel「前 10 項」一致）

    # ── 各篩選表門檻（自原檔樞紐 XML 還原）
    repeat_fault_min: int = 2           # 重複性故障：同設備同故障代碼 ≥ N 次
    repeat_fault_month: str | None = None   # None ＝ 期間最後一個月（原檔鎖「4月」）
    freq_repair_min: int = 10           # 設備故障維修 > N 次
    manhour_min: float = 15.0           # 維修工時篩選 ≥ N 小時
    recurrence_min: int = 5             # 故障原因再現性 ≥ N 次
    mttf_days_max: float = 30.0         # MTTF < N 天
    downtime_minutes_min: int = 1440    # 停機 > N 分
    repair_minutes_min: int = 720       # 維修 ≥ N 分
    availability_max: float = 0.99      # 妥善率 < N
    tiny: float = 0.01                  # 排除 0（未填寫）

    # ── 重要順序 → 允許結單天數（原檔 15 層 IF 攤平；不在表內者不判逾期）
    allow_days: dict = field(default_factory=lambda: {
        "緊急": 0, "一天": 1, "二天": 2, "三天": 3, "四天": 4, "五天": 5,
        "一週": 7, "兩週": 14, "三週": 21, "一月": 30, "一季": 90,
        "四個月": 120, "半年": 180, "一年": 365, "大修": 365,
    })


CFG = Config()

MONTHS = [f"{m}月" for m in range(1, 13)]
BLANK = "(空白)"

# 原檔樞紐的故障代碼欄序（Excel 中文筆畫序）。新代碼會接在後面。
FAULT_CODE_ORDER = [
    "其他製作安裝因素", "性能退化、老化", "負載變化過大", "異物入侵", "閒置過久",
    "過負載使用", "預防保養", "操作失誤", "機械疲勞", "磨耗、磨損、腐蝕",
    "環境惡劣(溫度、溼度、粉塵、振動...)", "過頻繁使用", "強度/容量不足",
    "接觸物腐蝕(化學物、液體、氣體...)",
]

REQUIRED = {
    "DATA": ["派工單號", "設備名稱", "組別", "維修狀態", "需求日", "派工日",
             "完成日", "重要順序"],
    "完工資料": ["派工單號", "設備名稱", "完成日", "故障原因", "故障代碼",
                 "維修類別", "停機時間(分)", "維修時間(分)"],
    "未完工資料": ["派工單號", "設備名稱", "異常名稱", "維修狀態", "組別", "需求日"],
}

DATA_COLS = ["派工單號", "請修單號", "設備名稱", "異常名稱", "故障描述", "組別",
             "維修人員", "維修狀態", "需求日", "派工日", "完成日", "重要順序",
             "結單日數", "符合度"]
KPI_COLS = ["派工單號", "設備編號", "設備名稱", "故障描述", "派工日", "完成日",
            "故障原因", "故障代碼", "完工報告", "維修類別", "停機時間(分)",
            "維修時間(分)", "Work Plan", "MTTF", "稼動率%", "妥善率%"]
OPEN_COLS = ["派工單號", "請修單號", "設備名稱", "異常名稱", "故障描述", "組別",
             "維修人員", "維修狀態", "需求日", "派工日", "完成日", "重要順序"]


# ══════════════════════════════════════════════════════════════════════════════
#  1. 載入、清理、重算
# ══════════════════════════════════════════════════════════════════════════════


class Log:
    """資料校正紀錄。每一筆 = 程式對輸入資料做了什麼、為什麼。"""

    def __init__(self):
        self.rows: list[dict] = []

    def add(self, sheet, key, item, before, after, reason):
        self.rows.append({"工作表": sheet, "派工單號／設備": key, "項目": item,
                          "輸入值": before, "程式值": after, "說明": reason})

    def frame(self) -> pd.DataFrame:
        cols = ["工作表", "派工單號／設備", "項目", "輸入值", "程式值", "說明"]
        return pd.DataFrame(self.rows, columns=cols)


def month_label(s: pd.Series) -> pd.Series:
    dt = pd.to_datetime(s, errors="coerce")
    return dt.dt.month.map(lambda m: f"{int(m)}月" if pd.notna(m) else None)


def load_sources(path: Path) -> dict[str, pd.DataFrame]:
    xl = pd.ExcelFile(path)
    missing = [s for s in REQUIRED if s not in xl.sheet_names]
    if missing:
        raise SystemExit(f"[錯誤] 來源檔缺少工作表：{missing}")
    out = {}
    for name, need in REQUIRED.items():
        df = xl.parse(name)
        df.columns = [str(c).strip() for c in df.columns]
        lack = [c for c in need if c not in df.columns]
        if lack:
            raise SystemExit(f"[錯誤] 工作表「{name}」缺少必要欄位：{lack}")
        out[name] = df.dropna(how="all").reset_index(drop=True)
    return out


def resolve_period(kpi: pd.DataFrame, cfg: Config) -> tuple[date, date]:
    if cfg.period_start and cfg.period_end:
        return cfg.period_start, cfg.period_end
    done = pd.to_datetime(kpi["完成日"], errors="coerce").dropna()
    if done.empty:
        raise SystemExit("[錯誤] 完工資料沒有可用的完成日，無法推算報表期間；請用 --period 指定")
    year = int(done.dt.year.mode().iloc[0])
    in_year = done[done.dt.year == year]
    last_m = int(in_year.dt.month.max())
    start = cfg.period_start or date(year, 1, 1)
    end = cfg.period_end or date(year, last_m, calendar.monthrange(year, last_m)[1])
    return start, end


def prepare(src: dict[str, pd.DataFrame], cfg: Config, log: Log) -> dict:
    # ── DATA ─────────────────────────────────────────────────────────────
    data = src["DATA"].copy()
    raw_n = len(data)
    data = data.drop_duplicates().reset_index(drop=True)
    if raw_n - len(data):
        log.add("DATA", "—", "完全重複列", raw_n, len(data),
                f"移除 {raw_n - len(data)} 列完全重複的匯出資料")
    dup_ids = data[data.duplicated("派工單號", keep=False)]["派工單號"].unique()
    if len(dup_ids):
        log.add("DATA", "、".join(map(str, dup_ids[:10])), "派工單號重複但內容不同",
                len(dup_ids), "保留全部", "同一單號多列且內容不同，請回 ERP 確認")
    for c in ("需求日", "派工日", "完成日"):
        data[c] = pd.to_datetime(data[c], errors="coerce")
    data["派工單號"] = data["派工單號"].astype(str).str.strip()

    asof = pd.Timestamp(cfg.as_of) if cfg.as_of else pd.Timestamp.today().normalize()
    calc_days = (data["完成日"].fillna(asof) - data["需求日"]).dt.days
    allow = data["重要順序"].map(cfg.allow_days)
    calc_fit = pd.Series("", index=data.index)
    calc_fit[(calc_days > allow) & allow.notna()] = "逾期"
    # 對帳：輸入檔若已有這兩欄，記錄差異（只記已完工單；未完工單本來就隨基準日變）
    if "結單日數" in data.columns:
        old = pd.to_numeric(data["結單日數"], errors="coerce")
        m = data["完成日"].notna() & old.notna() & ((old - calc_days).abs() > 0.5)
        for i in data.index[m][:200]:
            log.add("DATA", data.at[i, "派工單號"], "結單日數", old[i], int(calc_days[i]),
                    "完成日－需求日 重算")
    if "符合度" in data.columns:
        old = data["符合度"].fillna("").astype(str).str.strip()
        m = data["完成日"].notna() & (old != calc_fit)
        for i in data.index[m][:200]:
            log.add("DATA", data.at[i, "派工單號"], "符合度", old[i] or "(空白)",
                    calc_fit[i] or "(空白)", "依重要順序允許天數重判")
    data["結單日數"] = calc_days
    data["符合度"] = calc_fit
    data["需求月"] = month_label(data["需求日"])
    data["完成月"] = month_label(data["完成日"])

    # ── 完工資料 ──────────────────────────────────────────────────────────
    kpi = src["完工資料"].copy()
    kpi["派工單號"] = kpi["派工單號"].astype(str).str.strip()
    n0 = len(kpi)
    kpi = kpi.drop_duplicates("派工單號", keep="first").reset_index(drop=True)
    if n0 - len(kpi):
        log.add("完工資料", "—", "派工單號重複", n0, len(kpi), "同一派工單號只保留第一筆")
    for c in ("派工日", "完成日"):
        kpi[c] = pd.to_datetime(kpi[c], errors="coerce")
    for c in ("停機時間(分)", "維修時間(分)", "Work Plan", "MTTF", "稼動率%", "妥善率%"):
        if c not in kpi.columns:
            kpi[c] = pd.NA
        kpi[c] = pd.to_numeric(kpi[c], errors="coerce")
    for c in ("設備編號", "故障描述", "完工報告"):
        if c not in kpi.columns:
            kpi[c] = ""
    kpi["維修類別"] = kpi["維修類別"].astype(str).str.strip().str.upper()
    blank_cat = kpi["維修類別"].isin(["", "NAN", "NONE"])
    for i in kpi.index[blank_cat]:
        log.add("完工資料", kpi.at[i, "派工單號"], "維修類別", "(空白)", BLANK,
                "未填維修類別，不會進入 BD 類報表")
    kpi.loc[blank_cat, "維修類別"] = BLANK
    for c in ("故障代碼", "故障原因", "設備名稱"):
        kpi[c] = kpi[c].fillna("").astype(str).str.strip().replace("", BLANK)
    for c in ("停機時間(分)", "維修時間(分)"):
        bad = kpi[c].isna() | (kpi[c] < 0)
        for i in kpi.index[bad]:
            log.add("完工資料", kpi.at[i, "派工單號"], c, kpi.at[i, c], 0, "空白或負數視為 0")
        kpi.loc[bad, c] = 0
    kpi["Work Plan"] = kpi["Work Plan"].fillna(0)

    period_start, period_end = resolve_period(kpi, cfg)
    ps, pe = pd.Timestamp(period_start), pd.Timestamp(period_end)
    out_of_period = kpi["完成日"].isna() | (kpi["完成日"] < ps) | (kpi["完成日"] > pe)
    for i in kpi.index[out_of_period]:
        log.add("完工資料", kpi.at[i, "派工單號"], "完成日",
                kpi.at[i, "完成日"].date() if pd.notna(kpi.at[i, "完成日"]) else "(空白)",
                "排除", f"不在報表期間 {period_start}～{period_end} 內，不納入計算")
    kpi = kpi[~out_of_period].reset_index(drop=True)

    not_in_data = ~kpi["派工單號"].isin(set(data["派工單號"]))
    for i in kpi.index[not_in_data]:
        log.add("完工資料", kpi.at[i, "派工單號"], "派工單號", "不在 DATA", "保留",
                "ERP 匯出找不到此單號，請確認是否手動新增")

    kpi["維修工時"] = kpi["維修時間(分)"] / 60
    kpi["停機時數"] = kpi["停機時間(分)"] / 60
    kpi["完成月"] = month_label(kpi["完成日"])
    kpi = kpi_per_equipment(kpi, period_start, period_end, log, cfg.kpi_category)

    # ── 未完工資料 ────────────────────────────────────────────────────────
    op = src["未完工資料"].copy()
    op["派工單號"] = op["派工單號"].astype(str).str.strip()
    n0 = len(op)
    op = op.drop_duplicates("派工單號", keep="first").reset_index(drop=True)
    if n0 - len(op):
        log.add("未完工資料", "—", "派工單號重複", n0, len(op), "只保留第一筆")
    for c in ("需求日", "派工日", "完成日"):
        if c in op.columns:
            op[c] = pd.to_datetime(op[c], errors="coerce")
    closed_in_data = data[data["完成日"].notna()].set_index("派工單號")["完成日"]
    done_ids = op["派工單號"].map(closed_in_data)
    for i in op.index[done_ids.notna()]:
        log.add("未完工資料", op.at[i, "派工單號"], "維修狀態", "未完工",
                "排除", f"DATA 顯示已於 {done_ids[i].date()} 完成，不應列在未完工")
    op = op[done_ids.isna()].reset_index(drop=True)
    missing_open = data[data["完成日"].isna() & ~data["派工單號"].isin(set(op["派工單號"]))]
    for _, r in missing_open.iterrows():
        log.add("未完工資料", r["派工單號"], "未列入", "—", "未列入（僅提示）",
                f"DATA 中此單完成日空白（需求日 {r['需求日'].date() if pd.notna(r['需求日']) else '?'}），"
                "但未出現在未完工資料；若確實未結案請補列")
    for c in ("異常名稱", "維修狀態", "設備名稱", "故障描述", "組別"):
        if c in op.columns:
            op[c] = op[c].fillna("").astype(str).str.strip().replace("", BLANK)
    op["需求月"] = month_label(op["需求日"])
    op["需求年"] = op["需求日"].dt.year

    return {"DATA": data, "KPI": kpi, "OPEN": op, "period": (period_start, period_end),
            "as_of": asof.date()}


def kpi_per_equipment(kpi: pd.DataFrame, start: date, end: date, log: Log,
                      bd_category: str = "BD") -> pd.DataFrame:
    """每台設備重算 MTTF、稼動率、妥善率（以程式為準），填在該設備最後一列。

        期間分鐘 P ＝ (期間天數) × 1440
        MTTF(分)  ＝ (P − Σ停機分[全部類別]) ÷ 該設備完工單數[全部類別]
        稼動率    ＝ (P − Σ停機分[全部類別]) ÷ P
        妥善率    ＝ (P − Σ停機分[僅 BD 故障維修]) ÷ P
                    （計畫性工作 IM/MR/SI/OH 的停機不算「不妥善」；Work Plan 欄不參與計算）
    三條式子皆以 2024 年 1～4 月原檔 185 台設備逐台驗證，全數吻合。
    """
    P = ((pd.Timestamp(end) - pd.Timestamp(start)).days + 1) * 1440
    key = kpi["設備編號"].astype(str).str.strip()
    key = key.where(~key.isin(["", "nan", "None"]), kpi["設備名稱"])
    kpi["_eq"] = key
    kpi["_bd_down"] = kpi["停機時間(分)"].where(kpi["維修類別"] == bd_category, 0)
    g = kpi.groupby("_eq", sort=False)
    agg = g.agg(n=("派工單號", "count"), down=("停機時間(分)", "sum"),
                bd_down=("_bd_down", "sum"), old_mttf=("MTTF", "max"),
                old_av=("稼動率%", "max"), old_ok=("妥善率%", "max"),
                name=("設備名稱", "last"))
    agg["mttf"] = (P - agg["down"]) / agg["n"]
    agg["av"] = ((P - agg["down"]) / P).round(4)
    agg["ok"] = ((P - agg["bd_down"]) / P).round(4)

    for eq, r in agg.iterrows():
        for col, old, new, tol in (("MTTF", r["old_mttf"], r["mttf"], 1.0),
                                   ("稼動率%", r["old_av"], r["av"], 0.00006),
                                   ("妥善率%", r["old_ok"], r["ok"], 0.00006)):
            if pd.notna(old) and abs(old - new) > tol:
                log.add("完工資料", f"{r['name']}（{eq}）", col, old, round(new, 4),
                        "依停機分鐘與報表期間重算，覆寫輸入值")
    last_idx = g.tail(1).index
    for c in ("MTTF", "稼動率%", "妥善率%"):
        kpi[c] = pd.NA
    m = agg.loc[kpi.loc[last_idx, "_eq"]]
    kpi.loc[last_idx, "MTTF"] = m["mttf"].values
    kpi.loc[last_idx, "稼動率%"] = m["av"].values
    kpi.loc[last_idx, "妥善率%"] = m["ok"].values
    # 每列都帶一份（供程式內部彙總；輸出到工作表時仍只寫最後一列）
    kpi["_mttf_days"] = kpi["_eq"].map(agg["mttf"]) / 1440
    kpi["_avail"] = kpi["_eq"].map(agg["ok"])
    kpi["_last"] = False
    kpi.loc[last_idx, "_last"] = True
    return kpi


# ══════════════════════════════════════════════════════════════════════════════
#  2. 共用彙總工具
# ══════════════════════════════════════════════════════════════════════════════


def months_in(period) -> list[str]:
    s, e = period
    return [f"{m}月" for m in range(s.month, e.month + 1)] if s.year == e.year else MONTHS


def bd(b: dict, cfg: Config) -> pd.DataFrame:
    return b["KPI"][b["KPI"]["維修類別"] == cfg.kpi_category]


def by_month(df: pd.DataFrame, value: str | None, agg: str, months: list[str]) -> pd.DataFrame:
    """設備 × 月 交叉表（含總計欄）。value=None → 計數。"""
    if df.empty:
        return pd.DataFrame(columns=months + ["總計"])
    if value is None:
        pt = df.pivot_table(index="設備名稱", columns="完成月", values="派工單號",
                            aggfunc="count", fill_value=0)
    else:
        pt = df.pivot_table(index="設備名稱", columns="完成月", values=value,
                            aggfunc=agg, fill_value=0)
    pt = pt.reindex(columns=months, fill_value=0)
    pt["總計"] = pt.sum(axis=1)
    return pt


def order_codes(codes) -> list[str]:
    known = [c for c in FAULT_CODE_ORDER if c in set(codes)]
    extra = sorted(c for c in set(codes) if c not in FAULT_CODE_ORDER)
    return known + extra


def order_cats(cats, cfg: Config) -> list[str]:
    known = [c for c in cfg.category_order if c in set(cats)]
    extra = sorted(c for c in set(cats) if c not in cfg.category_order)
    return known + extra


# ══════════════════════════════════════════════════════════════════════════════
#  3. 寫入工具（版面／樣式／圖表）
# ══════════════════════════════════════════════════════════════════════════════

# 範本的顏色來自「樞紐分析表樣式」（佈景主題 Office 2007）：
#   PivotStyleMedium7／Medium14 → accent6 橘 F79646
#   PivotStyleLight16            → accent2 紅 C0504D（框線式）
#   PivotStyleLight18            → accent4 紫 8064A2（框線式）
# 這裡用明確的儲存格底色／框線把同樣的外觀畫出來。
ACCENT = {"orange": "F79646", "red": "C0504D", "purple": "8064A2", "blue": "4F81BD"}


def _tint(hex6: str, t: float) -> str:
    r, g, b = (int(hex6[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{int(round(c + (255 - c) * t)):02X}" for c in (r, g, b))


# 每張分析表的樣式：(色系, 是否框線式 Light 樣式, 全框線 Medium14, 列條紋)
SHEET_STYLE = {
    "完工類別統計總表": ("orange", False, False, False),
    "完工故障原因分析總表": ("orange", False, False, False),
    "未完工統計(機)": ("orange", False, False, False),
    "未完工統計(電)": ("purple", True, False, False),
    "結單率": ("red", True, False, False),
    "逾期結單": ("orange", False, True, False),
    "重複性故障": ("red", True, False, False),
    "設備故障維修次數總表": ("orange", False, False, True),
    "設備故障維修>10次": ("orange", False, False, False),
    "設備故障維修工時總表": ("orange", False, False, False),
    "設備故障維修工時篩選": ("orange", False, False, False),
    "PLAN WORK總表": ("orange", False, True, False),
    "PLAN WORK篩選": ("orange", False, False, False),
    "故障原因再現性總表": ("orange", False, True, True),
    "故障原因再現性篩選": ("orange", False, False, False),
    "MTTF總表": ("orange", False, False, False),
    "MTTF<30天": ("orange", False, False, False),
    "停機>24h": ("orange", False, True, False),
    "維修>12h": ("orange", False, True, False),
    "妥善率總表": ("orange", False, False, False),
    "妥善率<99%": ("orange", False, False, False),
}
# 分頁標籤顏色（範本）
TAB_COLOR = {
    "逾期結單": "7030A0", "設備故障維修次數總表": "FF0000", "設備故障維修工時總表": "FFC000",
    "PLAN WORK總表": "00B050", "故障原因再現性總表": "C00000",
    "完工資料": "1F497D", "未完工資料": "1F497D", "DATA": "1F497D",
}
# 圖表尺寸（cm，自範本錨點換算）
CHART_SIZE = {
    ("完工類別統計總表", 0): (8.0, 6.7), ("完工類別統計總表", 1): (7.9, 6.7), ("完工類別統計總表", 2): (8.0, 6.7),
    ("完工故障原因分析總表", 0): (15.2, 8.4), ("完工故障原因分析總表", 1): (15.2, 8.4), ("完工故障原因分析總表", 2): (15.2, 8.4),
    ("重複性故障", 0): (18.2, 10.9), ("重複性故障", 1): (18.1, 10.9),
    ("設備故障維修>10次", 0): (20.8, 11.9), ("設備故障維修工時篩選", 0): (25.0, 11.9),
    ("PLAN WORK篩選", 0): (18.2, 11.8), ("故障原因再現性篩選", 0): (47.5, 15.4),
    ("MTTF<30天", 0): (13.6, 12.2), ("妥善率<99%", 0): (20.6, 11.8),
}
# 欄寬（範本）
COL_WIDTH = {
    "完工類別統計總表": {"A": 13.75, "B": 6.38, "C": 10.88, "D": 11.25, "E": 10.88, "F": 11.25, "G": 12.5, "H": 11.12, "I": 11.12, "J": 12.62, "K": 16.62, "L": 6.38, "M": 10.88, "N": 11.25, "O": 10.88, "P": 11.25, "Q": 10.88},
    "完工故障原因分析總表": {"A": 42.38, "B": 6.5, "C": 10.12, "D": 11.5, "E": 10.12, "F": 11.5, "G": 10.12},
    "未完工統計(機)": {"A": 18.5, "B": 15.62, "C": 16.0, "D": 17.62, "E": 21.25, "F": 11.12, "G": 14.88, "H": 11.12},
    "未完工統計(電)": {"A": 18.5, "B": 13.88, "C": 20.12, "D": 17.12, "E": 50.12, "F": 10.88},
    "結單率": {"A": 16.38, "B": 11.12},
    "逾期結單": {"A": 10.88, "B": 13.5, "C": 13.5, "D": 18.12, "E": 24.5, "F": 31.25, "G": 13.5, "H": 18.62},
    "重複性故障": {"A": 33.88, "B": 42.5, "C": 20.0, "D": 8.5, "E": 28.0, "F": 8.5, "G": 42.5, "H": 14.5},
    "設備故障維修次數總表": {"A": 36.25, "B": 10.5},
    "設備故障維修>10次": {"A": 26.25, "B": 10.75},
    "設備故障維修工時總表": {"A": 36.25, "B": 10.5},
    "設備故障維修工時篩選": {"A": 28.75, "B": 14.12},
    "PLAN WORK總表": {"A": 36.25, "B": 11.12, "C": 10.62, "D": 7.12, "E": 10.62, "F": 7.12, "G": 10.62, "H": 7.12, "I": 10.62, "J": 9.12, "K": 10.62, "L": 15.25, "M": 12.25},
    "PLAN WORK篩選": {"A": 26.5, "B": 10.5, "C": 10.62, "D": 9.88},
    "故障原因再現性總表": {"A": 37.88, "B": 10.62},
    "故障原因再現性篩選": {"A": 29.12, "B": 12.88, "C": 12.62},
    "MTTF總表": {"A": 36.25, "B": 8.12},
    "MTTF<30天": {"A": 29.0, "B": 8.25},
    "停機>24h": {"A": 19.62, "B": 24.75, "C": 21.88, "D": 13.12, "E": 10.62},
    "維修>12h": {"A": 18.5, "B": 24.88, "C": 19.12, "D": 10.25, "E": 9.5, "F": 10.25, "G": 8.5, "H": 9.88},
    "妥善率總表": {"A": 36.25, "B": 16.62},
    "妥善率<99%": {"A": 18.88, "B": 9.88, "C": 9.5, "D": 6.62, "E": 9.62},
    "完工資料": {"A": 16.5, "B": 18.25, "C": 25.75, "D": 35.0, "E": 12.88, "F": 12.88, "G": 24.12, "H": 20.12, "I": 28.38, "J": 10.12, "K": 15.12, "L": 12.12, "M": 12.12, "N": 11.75, "O": 10.5, "P": 10.5},
    "未完工資料": {"A": 16.75, "B": 16.12, "C": 36.88, "D": 12.62, "E": 49.88, "F": 8.12, "G": 11.12, "H": 12.62, "I": 12.75, "J": 14.12, "K": 8.88, "L": 11.12},
    "DATA": {"A": 16.75, "B": 16.12, "C": 33.62, "D": 12.62, "E": 38.88, "F": 8.12, "G": 11.12, "H": 15.12, "I": 12.75, "J": 14.75, "K": 12.75, "L": 11.12, "M": 13.88, "N": 8.88},
}
FONT_NAME = "新細明體"
FONT_SIZE = 12
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT = "0.00%"
NUM2 = "0.00"


class W:
    """把資料寫到固定位置，並套用該工作表的樞紐樣式外觀。"""

    def __init__(self, ws, font_name: str = FONT_NAME):
        self.ws = ws
        name = ws.title
        accent, light, full_border, striped = SHEET_STYLE.get(name, ("orange", False, False, False))
        self.accent = ACCENT[accent]
        self.light, self.full_border, self.striped = light, full_border, striped
        self.font_name = font_name
        self.HDR = None if light else PatternFill("solid", fgColor=self.accent)
        self.HDR_FONT = Font(name=font_name, size=FONT_SIZE, bold=True,
                             color=None if light else "FFFFFF")
        self.TOT = PatternFill("solid", fgColor=_tint(self.accent, 0.8))
        self.STRIPE = PatternFill("solid", fgColor=_tint(self.accent, 0.9))
        acc_side = Side(style="medium", color=self.accent)
        self.HDR_BORDER = Border(top=acc_side, bottom=acc_side)
        self.TOT_BORDER = Border(top=Side(style="double", color=self.accent))
        ws.sheet_format.defaultRowHeight = 17.1
        for col, wdt in COL_WIDTH.get(name, {}).items():
            ws.column_dimensions[col].width = wdt
        if name in TAB_COLOR:
            ws.sheet_properties.tabColor = TAB_COLOR[name]

    def put(self, r, c, v, fmt=None, bold=False, fill=None):
        if isinstance(v, float) and pd.isna(v):
            v = None
        if hasattr(v, "item"):
            v = v.item()
        if isinstance(v, pd.Timestamp):
            v = v.to_pydatetime()
        cell = self.ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=self.font_name, size=FONT_SIZE, bold=bold)
        if fmt and isinstance(v, (int, float)):
            cell.number_format = fmt
        if fill is not None:
            cell.fill = fill
            if fill is self.TOT and self.light:
                cell.border = self.TOT_BORDER
        return cell

    def row(self, r, values, c0=1, fmts=None, bold=False, fill=None):
        for i, v in enumerate(values):
            f = fmts[i] if fmts and i < len(fmts) else None
            self.put(r, c0 + i, v, f, bold, fill)

    def filter_line(self, r, label, value):
        self.put(r, 1, label, bold=True)
        self.put(r, 2, value)

    def header(self, r, values, c0=1):
        for i, v in enumerate(values):
            cell = self.put(r, c0 + i, v)
            cell.font = self.HDR_FONT
            if self.HDR is not None:
                cell.fill = self.HDR
            else:
                cell.border = self.HDR_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def widths(self, widths: dict):
        for col, w in widths.items():
            if col not in COL_WIDTH.get(self.ws.title, {}):
                self.ws.column_dimensions[col].width = w

    def borders(self, r1, r2, c1, c2):
        """資料區收尾：Medium14 全框線；條紋樣式套交錯底色。"""
        if self.full_border:
            for row in self.ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                for cell in row:
                    cell.border = BORDER
        if self.striped:
            for i, row in enumerate(self.ws.iter_rows(min_row=r1 + 1, max_row=r2, min_col=c1, max_col=c2)):
                if i % 2 == 1 and row[0].fill.fill_type is None:
                    for cell in row:
                        cell.fill = self.STRIPE


def _fmt_cols(header: list[str]) -> list[str | None]:
    out = []
    for h in header:
        h = str(h)
        if "%" in h or "率" in h:
            out.append(PCT)
        elif any(k in h for k in ("工時", "時數", "MTTF")):
            out.append(NUM2)
        else:
            out.append(None)
    return out


def _blank0(v):
    """計數型樞紐：0 顯示為空白。"""
    return None if (v == 0 or v is None) else v


def chart(kind, ws, anchor, title, cats: Reference, series: list[Reference],
          titles_from_data=True, width=None, height=None, labels=False):
    ch = {"pie3d": PieChart3D, "pie": PieChart, "radar": RadarChart,
          "bar": BarChart, "bar3d": BarChart3D}[kind]()
    idx = len(ws._charts)
    size = CHART_SIZE.get((ws.title, idx), (15, 7.5))
    width, height = width or size[0], height or size[1]
    if kind in ("bar", "bar3d"):
        ch.type = "col"
        ch.grouping = "clustered"
    if kind == "radar":
        ch.type = "marker"
    ch.title = title or None
    for s in series:
        ch.add_data(s, titles_from_data=titles_from_data)
    ch.set_categories(cats)
    ch.width, ch.height = width, height
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ch.dataLabels.showVal = False
    ws.add_chart(ch, anchor)


# ══════════════════════════════════════════════════════════════════════════════
#  4. 21 張分析表
# ══════════════════════════════════════════════════════════════════════════════


def sh_完工類別統計總表(wb, b, cfg):
    ws = wb.create_sheet("完工類別統計總表"); w = W(ws, font_name="微軟正黑體")
    kpi = b["KPI"]; months = months_in(b["period"])
    cats = order_cats(kpi["維修類別"].unique(), cfg)
    g = kpi.groupby("維修類別")
    tot_n, tot_h, tot_d = len(kpi), kpi["維修工時"].sum(), kpi["停機時數"].sum()

    hdr = ["維修類別", "次數", "次數%", " 維修工時", "工時%", " 停機時數", "停機時數%",
           "工單平均維修工時", "工單平均停機時數"]
    w.header(1, hdr)
    fm = [None, None, PCT, NUM2, PCT, NUM2, PCT, NUM2, NUM2]
    r = 2
    for c in cats:
        d = g.get_group(c)
        n, h, dn = len(d), d["維修工時"].sum(), d["停機時數"].sum()
        w.row(r, [c, n, n / tot_n, h, h / tot_h if tot_h else 0, dn,
                  dn / tot_d if tot_d else 0, h / n, dn / n], fmts=fm)
        r += 1
    w.row(r, ["總計", tot_n, 1, tot_h, 1, tot_d, 1, tot_h / tot_n if tot_n else 0,
              tot_d / tot_n if tot_n else 0], fmts=fm, bold=True, fill=w.TOT)
    last = r
    w.borders(1, last, 1, 9)

    # 右側：月 × 維修類別
    hdr2 = ["完成日", "維修類別", "次數", "次數%", " 維修工時", " 工時%", " 停機時數", " 時數%"]
    w.header(1, hdr2, c0=10)
    fm2 = [None, None, None, PCT, NUM2, PCT, NUM2, PCT]
    r = 2
    for m in months:
        blk = kpi[kpi["完成月"] == m]
        if blk.empty:
            continue
        first = True
        for c in order_cats(blk["維修類別"].unique(), cfg):
            d = blk[blk["維修類別"] == c]
            n, h, dn = len(d), d["維修工時"].sum(), d["停機時數"].sum()
            w.row(r, [m if first else None, c, n, n / tot_n, h, h / tot_h if tot_h else 0,
                      dn, dn / tot_d if tot_d else 0], c0=10, fmts=fm2)
            first = False; r += 1
        n, h, dn = len(blk), blk["維修工時"].sum(), blk["停機時數"].sum()
        w.row(r, [f"{m} 合計", None, n, n / tot_n, h, h / tot_h if tot_h else 0, dn,
                  dn / tot_d if tot_d else 0], c0=10, fmts=fm2, bold=True, fill=w.TOT)
        r += 1
    w.row(r, ["總計", None, tot_n, 1, tot_h, 1, tot_d, 1], c0=10, fmts=fm2, bold=True, fill=w.TOT)
    w.borders(1, r, 10, 17)
    w.widths({"A": 14, "B": 8, "C": 10, "D": 11, "E": 10, "F": 11, "G": 11, "H": 16, "I": 16,
              "J": 10, "K": 10, "L": 8, "M": 10, "N": 11, "O": 10, "P": 11, "Q": 10})

    n = len(cats)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=n + 1)
    top = n + 5          # 圖表放在總計列下方，類別再多也不會蓋到表格
    chart("pie3d", ws, f"A{top}", "次數", cats_ref,
          [Reference(ws, min_col=2, min_row=1, max_row=n + 1)], labels=True)
    chart("pie3d", ws, f"D{top}", "工時", cats_ref,
          [Reference(ws, min_col=5, min_row=1, max_row=n + 1)], labels=True)
    chart("pie3d", ws, f"A{top + 16}", "停機時數", cats_ref,
          [Reference(ws, min_col=7, min_row=1, max_row=n + 1)], labels=True)


def sh_完工故障原因分析總表(wb, b, cfg):
    ws = wb.create_sheet("完工故障原因分析總表"); w = W(ws)
    kpi = b["KPI"]
    codes = order_codes(kpi["故障代碼"].unique())
    g = kpi.groupby("故障代碼")
    tot_n, tot_h, tot_d = len(kpi), kpi["維修工時"].sum(), kpi["停機時數"].sum()
    w.put(2, 2, "數值", bold=True)
    w.header(3, ["列標籤", "次數", "次數%", " 維修工時", " 工時%", " 停機時數", " 時數%"])
    fm = [None, None, PCT, NUM2, PCT, NUM2, PCT]
    r = 4
    for c in codes:
        d = g.get_group(c)
        n, h, dn = len(d), d["維修工時"].sum(), d["停機時數"].sum()
        w.row(r, [c, n, n / tot_n, h, h / tot_h if tot_h else 0, dn, dn / tot_d if tot_d else 0], fmts=fm)
        r += 1
    w.row(r, ["總計", tot_n, 1, tot_h, 1, tot_d, 1], fmts=fm, bold=True, fill=w.TOT)
    w.borders(3, r, 1, 7)
    w.widths({"A": 42, "B": 8, "C": 10, "D": 11, "E": 10, "F": 11, "G": 10})
    n = len(codes)
    cats_ref = Reference(ws, min_col=1, min_row=4, max_row=n + 3)
    chart("radar", ws, "H2", "工單次數", cats_ref, [Reference(ws, min_col=2, min_row=3, max_row=n + 3)])
    chart("radar", ws, "H16", "工單工時", cats_ref, [Reference(ws, min_col=4, min_row=3, max_row=n + 3)])
    chart("radar", ws, "N2", "停機時數", cats_ref, [Reference(ws, min_col=6, min_row=3, max_row=n + 3)])


def _open_stat(wb, b, cfg, name: str, group: str, with_id: bool):
    ws = wb.create_sheet(name); w = W(ws)
    op = b["OPEN"]
    df = op[op["組別"] == group].copy()
    idx = ["異常名稱", "維修狀態", "設備名稱"] + (["派工單號"] if with_id else []) + ["故障描述"]
    w.filter_line(2, "組別", group)
    w.put(4, 1, "計數 - 派工單號", bold=True)
    if df.empty or df["需求日"].isna().all():
        w.header(6, idx + ["總計"])
        w.row(7, ["總計"] + [None] * (len(idx) - 1) + [0], bold=True, fill=w.TOT)
        return
    df = df[df["需求日"].notna()]
    df["需求年"] = df["需求年"].astype(int)
    years = sorted(df["需求年"].unique())
    cols = []   # (year, month or None(=合計))
    for y in years:
        for m in sorted(df[df["需求年"] == y]["需求日"].dt.month.unique()):
            cols.append((y, int(m)))
        cols.append((y, None))
    c0 = len(idx) + 1
    w.put(4, c0, "年", bold=True); w.put(4, c0 + 1, "需求日", bold=True)
    # 第 5 列：年標籤；第 6 列：月標籤
    for j, (y, m) in enumerate(cols):
        c = c0 + j
        if m is None:
            w.put(5, c, f"{y}年 合計", bold=True, fill=w.HDR)
        elif j == 0 or cols[j - 1][0] != y:
            w.put(5, c, f"{y}年", bold=True, fill=w.HDR)
        w.put(6, c, f"{m}月" if m else None, bold=True, fill=w.HDR)
    w.put(5, c0 + len(cols), "總計", bold=True, fill=w.HDR)
    w.header(6, idx)
    # 排序：狀態空白排最後，其餘依首次出現
    status_order = {s: i for i, s in enumerate([s for s in df["維修狀態"].unique() if s != BLANK])}
    df["_s"] = df["維修狀態"].map(lambda s: status_order.get(s, 999))
    df = df.sort_values(["異常名稱", "_s", "設備名稱", "需求日"]).reset_index(drop=True)
    r = 7
    prev = [None] * len(idx)
    col_tot = [0] * (len(cols) + 1)
    for _, rec in df.iterrows():
        vals = [rec[c] for c in idx]
        show = []
        same = True
        for k, v in enumerate(vals):
            if same and v == prev[k] and k < len(idx) - 1:
                show.append(None)
            else:
                same = False
                show.append(v)
        prev = vals
        w.row(r, show)
        y, m = rec["需求年"], rec["需求日"].month
        for j, (yy, mm) in enumerate(cols):
            if yy == y and (mm == m or mm is None):
                w.put(r, c0 + j, 1); col_tot[j] += 1
        w.put(r, c0 + len(cols), 1); col_tot[-1] += 1
        r += 1
    w.put(r, 1, "總計", bold=True, fill=w.TOT)
    for j, t in enumerate(col_tot):
        w.put(r, c0 + j, t, bold=True, fill=w.TOT)
    w.borders(6, r, 1, c0 + len(cols))
    w.widths({"A": 18, "B": 15, "C": 22, "D": 18, "E": 30 if with_id else 24})


def sh_未完工統計_機(wb, b, cfg):
    _open_stat(wb, b, cfg, "未完工統計(機)", "機械組", with_id=False)


def sh_未完工統計_電(wb, b, cfg):
    _open_stat(wb, b, cfg, "未完工統計(電)", "儀電組", with_id=True)


def sh_結單率(wb, b, cfg):
    ws = wb.create_sheet("結單率"); w = W(ws)
    d = b["DATA"]; year = b["period"][0].year
    o = d[d["需求日"].dt.year == year]["需求月"].value_counts()
    c = d[d["完成日"].dt.year == year]["完成月"].value_counts()
    months = [m for m in MONTHS if m in set(o.index) | set(c.index)]
    w.put(1, 2, "欄標籤", bold=True)
    w.header(2, months + ["總計"], c0=2)
    w.put(3, 1, "工單開立數量", bold=True)
    w.row(3, [int(o.get(m, 0)) for m in months] + [int(o.sum())], c0=2)
    w.put(4, 2, "欄標籤", bold=True)
    w.header(5, months + ["總計"], c0=2)
    w.put(6, 1, "工單完工數量", bold=True)
    w.row(6, [int(c.get(m, 0)) for m in months] + [int(c.sum())], c0=2)
    w.put(7, 1, "結單率", bold=True)
    for j in range(len(months) + 1):
        col = get_column_letter(2 + j)
        w.put(7, 2 + j, f"=IF({col}3=0,\"\",{col}6/{col}3)", fmt=PCT)
        ws.cell(7, 2 + j).number_format = PCT
    w.borders(2, 7, 1, len(months) + 2)
    w.widths({"A": 16})


def sh_逾期結單(wb, b, cfg):
    ws = wb.create_sheet("逾期結單"); w = W(ws)
    d = b["DATA"]; ps = pd.Timestamp(b["period"][0])
    df = d[(d["需求日"] >= ps) & (d["符合度"] == "逾期")].copy()
    w.filter_line(2, "符合度", "逾期")
    hdr = ["需求日", "維修狀態", "重要順序", "派工單號", "設備名稱", "故障描述", "結單日數", "計數 - 派工單號"]
    w.header(4, hdr)
    df["維修狀態"] = df["維修狀態"].fillna(BLANK)
    df["_m"] = df["需求日"].dt.month
    df["_p"] = df["重要順序"].map(cfg.allow_days).fillna(999)
    df = df.sort_values(["_m", "維修狀態", "_p", "派工單號"], ascending=[True, True, True, False])
    r = 5
    grand = 0
    for m, blk in df.groupby("_m", sort=True):
        prev = (None, None)
        first = True
        for _, rec in blk.iterrows():
            key = (rec["維修狀態"], rec["重要順序"])
            w.row(r, [f"{m}月" if first else None,
                      rec["維修狀態"] if key[0] != prev[0] else None,
                      rec["重要順序"] if key != prev else None,
                      rec["派工單號"], rec["設備名稱"], rec.get("故障描述"),
                      int(rec["結單日數"]), 1])
            prev = key; first = False; r += 1
        w.row(r, [f"{m}月 合計"] + [None] * 6 + [len(blk)], bold=True, fill=w.TOT)
        grand += len(blk); r += 1
    w.row(r, ["總計"] + [None] * 6 + [grand], bold=True, fill=w.TOT)
    w.borders(4, r, 1, 8)
    w.widths({"A": 11, "B": 13, "C": 10, "D": 16, "E": 26, "F": 40, "G": 9, "H": 14})


def sh_重複性故障(wb, b, cfg):
    ws = wb.create_sheet("重複性故障"); w = W(ws)
    months = months_in(b["period"])
    month = cfg.repeat_fault_month or months[-1]
    df = bd(b, cfg)
    df = df[df["完成月"] == month]
    g = (df.groupby(["設備名稱", "故障代碼"])["派工單號"].count()
           .reset_index(name="次數"))
    keep = g[g["次數"] >= cfg.repeat_fault_min].sort_values(["設備名稱"]).reset_index(drop=True)
    w.filter_line(1, "維修類別", cfg.kpi_category)
    w.filter_line(2, "完成日", month)
    w.header(4, ["設備名稱", "故障代碼", "計數 - 派工單號"])
    w.header(4, ["故障代碼", "次數"], c0=5)
    w.header(4, ["列標籤", "加總 - 次數"], c0=7)
    r = 5
    for _, rec in keep.iterrows():
        w.row(r, [rec["設備名稱"], rec["故障代碼"], int(rec["次數"])])
        w.row(r, [rec["故障代碼"], int(rec["次數"])], c0=5)
        r += 1
    w.row(r, ["總計", None, int(keep["次數"].sum())], bold=True, fill=w.TOT)
    n = len(keep)
    byc = keep.groupby("故障代碼")["次數"].sum()
    byc = byc.reindex(order_codes(byc.index))
    rr = 5
    for code, v in byc.items():
        w.row(rr, [code, int(v)], c0=7); rr += 1
    w.row(rr, ["總計", int(byc.sum())], c0=7, bold=True, fill=w.TOT)
    w.borders(4, r, 1, 3); w.borders(4, max(rr, 5), 5, 8)
    w.widths({"A": 34, "B": 42, "C": 16, "E": 42, "F": 8, "G": 42, "H": 12})
    if n:
        chart("pie", ws, "A17", "重複性故障（依設備）",
              Reference(ws, min_col=1, min_row=5, max_row=4 + n),
              [Reference(ws, min_col=3, min_row=4, max_row=4 + n)], labels=True)
        chart("pie", ws, "D17", "故障原因比例",
              Reference(ws, min_col=7, min_row=5, max_row=4 + len(byc)),
              [Reference(ws, min_col=8, min_row=4, max_row=4 + len(byc))], labels=True)


def _equip_month_sheet(wb, b, cfg, name, value, agg, label, filt=None,
                       total_row=True, blank_zero=False, sort="asc",
                       chart_anchor=None, chart_kind="bar", chart_title=None, filter_val=None):
    ws = wb.create_sheet(name); w = W(ws)
    months = months_in(b["period"])
    pt = by_month(bd(b, cfg), value, agg, months)
    if filt is not None and not pt.empty:
        pt = pt[filt(pt["總計"])]
    if sort == "asc":
        pt = pt.sort_values("總計", kind="stable")
    elif sort == "desc":
        pt = pt.sort_values("總計", ascending=False, kind="stable")
    w.filter_line(1, "維修類別", filter_val or cfg.kpi_category)
    w.put(3, 1, label, bold=True); w.put(3, 2, "欄標籤", bold=True)
    hdr = ["列標籤"] + months + ["總計"]
    w.header(4, hdr)
    fm = [None] + [NUM2 if value else None] * (len(months) + 1)
    r = 5
    for eq, rec in pt.iterrows():
        vals = [rec[m] for m in months] + [rec["總計"]]
        if blank_zero:
            vals = [_blank0(v) for v in vals]
        w.row(r, [eq] + vals, fmts=fm); r += 1
    if total_row:
        tot = [pt[m].sum() for m in months] + [pt["總計"].sum()] if not pt.empty else [0] * (len(months) + 1)
        w.row(r, ["總計"] + tot, fmts=fm, bold=True, fill=w.TOT)
    w.borders(4, r, 1, len(hdr))
    w.widths({"A": 36, "B": 8})
    n = len(pt)
    if chart_anchor and n:
        chart(chart_kind, ws, chart_anchor, chart_title,
              Reference(ws, min_col=1, min_row=5, max_row=4 + n),
              [Reference(ws, min_col=2 + j, min_row=4, max_row=4 + n) for j in range(len(months))])
    return pt


def sh_設備故障維修次數總表(wb, b, cfg):
    _equip_month_sheet(wb, b, cfg, "設備故障維修次數總表", None, "count", "計數 - 派工單號",
                       blank_zero=True)


def sh_設備故障維修_高次數(wb, b, cfg):
    _equip_month_sheet(wb, b, cfg, "設備故障維修>10次", None, "count", "計數 - 派工單號",
                       filt=lambda t: t > cfg.freq_repair_min, blank_zero=True,
                       chart_anchor="N3", chart_title=f"故障次數>{cfg.freq_repair_min}")


def sh_設備故障維修工時總表(wb, b, cfg):
    _equip_month_sheet(wb, b, cfg, "設備故障維修工時總表", "維修工時", "sum", "加總 - 維修工時")


def sh_設備故障維修工時篩選(wb, b, cfg):
    _equip_month_sheet(wb, b, cfg, "設備故障維修工時篩選", "維修工時", "sum", "加總 - 維修工時",
                       filt=lambda t: t >= cfg.manhour_min,
                       chart_anchor="P5", chart_title=f"工時總計>={cfg.manhour_min:g}小時")


def sh_PLAN_WORK總表(wb, b, cfg):
    ws = wb.create_sheet("PLAN WORK總表"); w = W(ws)
    kpi = b["KPI"]
    cats = order_cats(kpi["維修類別"].unique(), cfg)
    pt = kpi.pivot_table(index="設備名稱", columns="維修類別", values="維修工時",
                         aggfunc="sum", fill_value=0).reindex(columns=cats, fill_value=0)
    pt["_tot"] = pt.sum(axis=1)
    pt = pt.sort_values("_tot", kind="stable")
    grand = pt["_tot"].sum()
    w.put(1, 2, "欄標籤", bold=True)
    for j, c in enumerate(cats):
        w.put(2, 2 + 2 * j, c, bold=True, fill=w.HDR)
    last_c = 2 + 2 * len(cats)
    w.put(2, last_c, " 工時 的加總", bold=True, fill=w.HDR)
    w.put(2, last_c + 1, " % 的加總", bold=True, fill=w.HDR)
    hdr = ["列標籤"] + [" 工時", " %"] * len(cats) + [None, None]
    w.header(3, hdr)
    fm = [None] + [NUM2, PCT] * len(cats) + [NUM2, PCT]
    r = 4
    for eq, rec in pt.iterrows():
        vals = []
        for c in cats:
            ct = pt[c].sum()
            vals += [rec[c], rec[c] / ct if ct else 0]
        vals += [rec["_tot"], rec["_tot"] / grand if grand else 0]
        w.row(r, [eq] + vals, fmts=fm); r += 1
    tot = []
    for c in cats:
        tot += [pt[c].sum(), 1]
    tot += [grand, 1]
    w.row(r, ["總計"] + tot, fmts=fm, bold=True, fill=w.TOT)
    w.borders(2, r, 1, last_c + 1)
    w.widths({"A": 36})


def sh_PLAN_WORK篩選(wb, b, cfg):
    ws = wb.create_sheet("PLAN WORK篩選"); w = W(ws)
    kpi = b["KPI"]
    df = kpi[kpi["維修類別"].isin(cfg.planwork_categories)]
    cats = [c for c in cfg.planwork_categories if c in set(df["維修類別"])]
    w.put(1, 1, "加總 - 維修工時", bold=True); w.put(1, 2, "欄標籤", bold=True)
    hdr = ["列標籤"] + cats + ["總計"]
    w.header(2, hdr)
    if df.empty:
        w.row(3, ["總計"] + [None] * len(cats) + [0], bold=True, fill=w.TOT)
        return
    pt = df.pivot_table(index="設備名稱", columns="維修類別", values="維修工時",
                        aggfunc="sum", fill_value=0).reindex(columns=cats, fill_value=0)
    pt["_tot"] = pt.sum(axis=1)
    # 前 N 名（含同分）：Excel「前 10 項」遇同值會一併列出
    thr = pt["_tot"].sort_values(ascending=False).head(cfg.planwork_top_n).min()
    top = pt[pt["_tot"] >= thr]
    fm = [None] + [PCT] * (len(cats) + 1)
    r = 3
    for eq, rec in top.iterrows():
        vals = [rec[c] / pt[c].sum() if pt[c].sum() else 0 for c in cats]
        vals.append(rec["_tot"] / pt["_tot"].sum())
        w.row(r, [eq] + vals, fmts=fm); r += 1
    w.row(r, ["總計"] + [1] * (len(cats) + 1), fmts=fm, bold=True, fill=w.TOT)
    w.borders(2, r, 1, len(hdr))
    w.widths({"A": 28, "B": 10, "C": 10, "D": 10})
    n = len(top)
    chart("bar", ws, "F1", f"TOP {cfg.planwork_top_n}",
          Reference(ws, min_col=1, min_row=3, max_row=2 + n),
          [Reference(ws, min_col=2 + j, min_row=2, max_row=2 + n) for j in range(len(cats))])


def _recurrence(wb, b, cfg, name, min_total=None, chart_anchor=None):
    ws = wb.create_sheet(name); w = W(ws)
    df = bd(b, cfg)
    codes = order_codes(df["故障代碼"].unique())
    w.filter_line(1, "維修類別", cfg.kpi_category)
    w.put(3, 1, "計數 - 故障原因", bold=True); w.put(3, 2, "欄標籤", bold=True)
    hdr = ["列標籤"] + codes + ["總計"]
    w.header(4, hdr)
    if df.empty:
        return
    # 列序：依設備編號首次出現（原檔完工資料即以設備編號排序）
    order = list(dict.fromkeys(df["設備名稱"]))
    pt = df.pivot_table(index="設備名稱", columns="故障代碼", values="派工單號",
                        aggfunc="count", fill_value=0).reindex(columns=codes, fill_value=0)
    pt["總計"] = pt.sum(axis=1)
    pt = pt.loc[[o for o in order if o in pt.index]]
    if min_total is not None:
        pt = pt[pt["總計"] >= min_total]
    r = 5
    for eq, rec in pt.iterrows():
        w.row(r, [eq] + [_blank0(rec[c]) for c in codes] + [rec["總計"]]); r += 1
    w.row(r, ["總計"] + [_blank0(pt[c].sum()) for c in codes] + [pt["總計"].sum()],
          bold=True, fill=w.TOT)
    w.borders(4, r, 1, len(hdr))
    w.widths({"A": 36})
    n = len(pt)
    if chart_anchor and n:
        chart("bar", ws, chart_anchor, f"故障原因>={min_total}種",
              Reference(ws, min_col=1, min_row=5, max_row=4 + n),
              [Reference(ws, min_col=2 + j, min_row=4, max_row=4 + n) for j in range(len(codes))])


def sh_故障原因再現性總表(wb, b, cfg):
    _recurrence(wb, b, cfg, "故障原因再現性總表")


def sh_故障原因再現性篩選(wb, b, cfg):
    _recurrence(wb, b, cfg, "故障原因再現性篩選", cfg.recurrence_min, chart_anchor="A31")


def _equip_kpi(df: pd.DataFrame, col: str) -> pd.Series:
    """每台設備一個值（程式已對每列帶同一值，取 max 即可）。"""
    return df.groupby("設備名稱")[col].max()


def sh_MTTF總表(wb, b, cfg):
    ws = wb.create_sheet("MTTF總表"); w = W(ws)
    s = _equip_kpi(bd(b, cfg), "_mttf_days").sort_values(ascending=False)
    w.filter_line(1, "維修類別", cfg.kpi_category)
    w.header(3, ["列標籤", " MTTF"])
    r = 4
    for eq, v in s.items():
        w.row(r, [eq, round(float(v), 4)], fmts=[None, NUM2]); r += 1
    w.borders(3, max(r - 1, 3), 1, 2)
    w.widths({"A": 36, "B": 10})


def _kpi_by_last_month(wb, b, cfg, name, col, label, lo, hi, sort_desc, chart_anchor, chart_title, fmt):
    ws = wb.create_sheet(name); w = W(ws)
    months = months_in(b["period"])
    df = bd(b, cfg)
    w.filter_line(1, "維修類別", cfg.kpi_category)
    w.put(3, 1, label, bold=True); w.put(3, 2, "完成日" if col == "_mttf_days" else "欄標籤", bold=True)
    hdr = (["設備名稱"] if col == "_mttf_days" else ["列標籤"]) + months + (["總計"] if col == "_mttf_days" else [])
    w.header(4, hdr)
    if df.empty:
        return
    last_m = df.sort_values("完成日").groupby("設備名稱")["完成月"].last()
    val = _equip_kpi(df, col)
    keep = val[(val >= lo) & (val < hi)].sort_values(ascending=not sort_desc)
    r = 5
    for eq, v in keep.items():
        vals = [round(float(v), 4) if m == last_m[eq] else 0 for m in months]
        if col == "_mttf_days":
            vals.append(round(float(v), 4))
        w.row(r, [eq] + vals, fmts=[None] + [fmt] * len(vals)); r += 1
    w.borders(4, max(r - 1, 4), 1, len(hdr))
    w.widths({"A": 30})
    n = len(keep)
    if n:
        chart("bar3d", ws, chart_anchor, chart_title,
              Reference(ws, min_col=1, min_row=5, max_row=4 + n),
              [Reference(ws, min_col=2 + j, min_row=4, max_row=4 + n) for j in range(len(months))])


def sh_MTTF_低於門檻(wb, b, cfg):
    _kpi_by_last_month(wb, b, cfg, "MTTF<30天", "_mttf_days", " MTTF", cfg.tiny, cfg.mttf_days_max,
                       True, "O4", f"MTTF<{cfg.mttf_days_max:g}天之設備項目", NUM2)


def _order_detail(wb, b, cfg, name, minute_col, hour_col, label, cond):
    ws = wb.create_sheet(name); w = W(ws)
    df = bd(b, cfg)
    df = df[cond(df[minute_col])].sort_values([minute_col, "派工單號"])
    months = [m for m in months_in(b["period"]) if m in set(df["完成月"])]
    w.filter_line(1, "維修類別", cfg.kpi_category)
    w.put(3, 1, label, bold=True); w.put(3, 5, "完成日", bold=True)
    hdr = ["派工單號", "設備名稱", "故障原因", minute_col] + months + ["總計"]
    w.header(4, hdr)
    r = 5
    for _, rec in df.iterrows():
        vals = [rec[hour_col] if rec["完成月"] == m else 0 for m in months]
        w.row(r, [rec["派工單號"], rec["設備名稱"], rec["故障原因"], rec[minute_col]] + vals + [rec[hour_col]],
              fmts=[None] * 4 + [NUM2] * (len(months) + 1)); r += 1
    tot = [df[df["完成月"] == m][hour_col].sum() for m in months] + [df[hour_col].sum()]
    w.row(r, ["總計", None, None, None] + tot, fmts=[None] * 4 + [NUM2] * (len(months) + 1),
          bold=True, fill=w.TOT)
    w.borders(4, r, 1, len(hdr))
    w.widths({"A": 18, "B": 26, "C": 30, "D": 12})


def sh_停機超時(wb, b, cfg):
    _order_detail(wb, b, cfg, "停機>24h", "停機時間(分)", "停機時數", "加總 - 停機時數",
                  lambda s: s > cfg.downtime_minutes_min)


def sh_維修超時(wb, b, cfg):
    _order_detail(wb, b, cfg, "維修>12h", "維修時間(分)", "維修工時", "加總 - 維修工時",
                  lambda s: s >= cfg.repair_minutes_min)


def sh_妥善率總表(wb, b, cfg):
    ws = wb.create_sheet("妥善率總表"); w = W(ws)
    s = _equip_kpi(b["KPI"], "_avail").sort_values(ascending=False)
    w.filter_line(1, "維修類別", "(全部)")
    w.header(3, ["列標籤", "最大 - 妥善率%"])
    r = 4
    for eq, v in s.items():
        w.row(r, [eq, round(float(v), 4)], fmts=[None, PCT]); r += 1
    w.borders(3, max(r - 1, 3), 1, 2)
    w.widths({"A": 36, "B": 16})


def sh_妥善率_低於門檻(wb, b, cfg):
    _kpi_by_last_month(wb, b, cfg, "妥善率<99%", "_avail", "最大 - 妥善率%", cfg.tiny,
                       cfg.availability_max, False, "N3",
                       f"妥善率<{cfg.availability_max:.0%}之設備項目", PCT)


# ══════════════════════════════════════════════════════════════════════════════
#  5. 三張來源表（清理＋重算後）與校正紀錄
# ══════════════════════════════════════════════════════════════════════════════


def _write_table(wb, name: str, df: pd.DataFrame, cols: list[str], table_name: str,
                 date_cols=(), fmts: dict | None = None):
    ws = wb.create_sheet(name); w = W(ws)
    use = [c for c in cols if c in df.columns]
    w.HDR = PatternFill("solid", fgColor=ACCENT["blue"]); w.header(1, use)
    for r, (_, rec) in enumerate(df[use].iterrows(), start=2):
        for c, col in enumerate(use, start=1):
            v = rec[col]
            if pd.isna(v) if not isinstance(v, str) else False:
                v = None
            cell = w.put(r, c, v)
            if col in date_cols and v is not None:
                cell.number_format = "yyyy/m/d"
            elif fmts and col in fmts and isinstance(v, (int, float)):
                cell.number_format = fmts[col]
    ref = f"A1:{get_column_letter(len(use))}{max(len(df) + 1, 2)}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "A2"


def sh_sources(wb, b):
    kpi = b["KPI"].copy()
    for c in ("MTTF", "稼動率%", "妥善率%"):
        kpi.loc[~kpi["_last"], c] = pd.NA
    kpi["Work Plan"] = kpi["Work Plan"].where(kpi["Work Plan"] != 0, pd.NA)
    _write_table(wb, "完工資料", kpi, KPI_COLS, "KPI_data", ("派工日", "完成日"),
                 {"稼動率%": "0.0000", "妥善率%": "0.0000", "MTTF": "0.00"})
    _write_table(wb, "未完工資料", b["OPEN"], OPEN_COLS, "未完工資料", ("需求日", "派工日", "完成日"))
    data = b["DATA"].copy()
    data["符合度"] = data["符合度"].replace("", pd.NA)
    _write_table(wb, "DATA", data, DATA_COLS, "DATA", ("需求日", "派工日", "完成日"))


def sh_log(wb, b, log: Log, cfg: Config):
    ws = wb.create_sheet("_資料校正紀錄"); w = W(ws)
    s, e = b["period"]
    P = ((pd.Timestamp(e) - pd.Timestamp(s)).days + 1)
    info = [
        ("報表期間", f"{s} ～ {e}（{P} 天，{P * 1440:,} 分鐘）"),
        ("未完工單結單日數基準日", str(b["as_of"])),
        ("KPI 維修類別", cfg.kpi_category),
        ("MTTF(分)", "(期間分鐘 − Σ停機分[全部類別]) ÷ 該設備完工單數"),
        ("稼動率%", "(期間分鐘 − Σ停機分[全部類別]) ÷ 期間分鐘"),
        ("妥善率%", f"(期間分鐘 − Σ停機分[僅 {cfg.kpi_category}]) ÷ 期間分鐘（Work Plan 不參與）"),
        ("結單日數", "完成日 − 需求日（未完工：基準日 − 需求日）"),
        ("符合度", "結單日數 > 重要順序允許天數 → 逾期"),
        ("原則", "以上欄位一律由程式重算；輸入檔既有值只用來對帳，差異列於下表"),
    ]
    w.header(1, ["參數", "值"])
    for i, (k, v) in enumerate(info, start=2):
        w.row(i, [k, v])
    r0 = len(info) + 3
    df = log.frame()
    w.header(r0, list(df.columns))
    for i, (_, rec) in enumerate(df.iterrows(), start=r0 + 1):
        w.row(i, [None if (not isinstance(v, str) and pd.isna(v)) else v for v in rec])
    w.widths({"A": 26, "B": 34, "C": 16, "D": 16, "E": 16, "F": 70})
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)


# ══════════════════════════════════════════════════════════════════════════════
#  6. 主流程
# ══════════════════════════════════════════════════════════════════════════════

BUILDERS = [
    sh_完工類別統計總表, sh_完工故障原因分析總表, sh_未完工統計_機, sh_未完工統計_電,
    sh_結單率, sh_逾期結單, sh_重複性故障, sh_設備故障維修次數總表, sh_設備故障維修_高次數,
    sh_設備故障維修工時總表, sh_設備故障維修工時篩選, sh_PLAN_WORK總表, sh_PLAN_WORK篩選,
    sh_故障原因再現性總表, sh_故障原因再現性篩選, sh_MTTF總表, sh_MTTF_低於門檻,
    sh_停機超時, sh_維修超時, sh_妥善率總表, sh_妥善率_低於門檻,
]


def build(src_path: Path, out_path: Path, cfg: Config = CFG) -> dict:
    log = Log()
    src = load_sources(src_path)
    b = prepare(src, cfg, log)
    wb = Workbook()
    wb.remove(wb.active)
    failed = []
    for fn in BUILDERS:
        try:
            fn(wb, b, cfg)
        except Exception as e:  # 單張表失敗不拖垮整份
            name = fn.__name__.replace("sh_", "")
            failed.append((name, f"{type(e).__name__}: {e}"))
            ws = wb.create_sheet(name[:31])
            ws["A1"] = f"產表失敗：{type(e).__name__}: {e}"
    sh_sources(wb, b)
    sh_log(wb, b, log, cfg)
    wb.save(out_path)
    return {"period": b["period"], "as_of": b["as_of"], "log": log.frame(),
            "rows": {"DATA": len(b["DATA"]), "完工資料": len(b["KPI"]), "未完工資料": len(b["OPEN"])},
            "failed": failed}


def check(src_path: Path) -> dict:
    """只檢查格式，不產表。回傳可直接判讀的字典（--check 會印成 JSON）。"""
    import json  # noqa: F401  (供呼叫端使用)
    xl = pd.ExcelFile(src_path)
    rep = {"file": str(src_path), "sheets_found": xl.sheet_names, "ok": True,
           "missing_sheets": [], "missing_columns": {}, "rows": {}, "period_guess": None,
           "notes": []}
    for name, need in REQUIRED.items():
        if name not in xl.sheet_names:
            rep["missing_sheets"].append(name); rep["ok"] = False
            continue
        df = xl.parse(name).dropna(how="all")
        df.columns = [str(c).strip() for c in df.columns]
        lack = [c for c in need if c not in df.columns]
        if lack:
            rep["missing_columns"][name] = lack; rep["ok"] = False
        rep["rows"][name] = int(len(df))
        if name == "完工資料" and not lack:
            done = pd.to_datetime(df["完成日"], errors="coerce").dropna()
            if len(done):
                y = int(done.dt.year.mode().iloc[0])
                rep["period_guess"] = f"{y}-01-01 ~ {y}-{int(done[done.dt.year == y].dt.month.max()):02d} 月底"
            else:
                rep["notes"].append("完工資料的完成日全部空白，需用 --period 指定期間")
            if len(df) == 0:
                rep["notes"].append("完工資料沒有任何資料列")
    if "DATA" in rep["rows"] and rep["rows"]["DATA"] == 0:
        rep["notes"].append("DATA 沒有任何資料列")
    return rep


def main():
    ap = argparse.ArgumentParser(description="從 DATA／完工資料／未完工資料 產出 ERP 維護績效月報（範本格式）")
    ap.add_argument("source", type=Path, help="來源 xlsx（含三張工作表）")
    ap.add_argument("--check", action="store_true", help="只檢查格式並印出 JSON 摘要，不產表")
    ap.add_argument("-o", "--output", type=Path, default=None)
    ap.add_argument("--period", nargs=2, metavar=("起", "迄"), default=None,
                    help="報表期間 YYYY-MM-DD YYYY-MM-DD（省略＝依完工資料自動推算）")
    ap.add_argument("--as-of", default=None, help="未完工單結單日數基準日（省略＝今天）")
    ap.add_argument("--category", default=CFG.kpi_category, help="KPI 維修類別（預設 BD）")
    ap.add_argument("--repeat-month", default=None, help="重複性故障看哪個月，如 4月（省略＝期間最後一月）")
    a = ap.parse_args()

    if a.check:
        import json
        rep = check(a.source)
        print(json.dumps(rep, ensure_ascii=False, indent=2))
        sys.exit(0 if rep["ok"] else 2)

    cfg = Config(kpi_category=a.category, repeat_fault_month=a.repeat_month)
    if a.period:
        cfg.period_start = datetime.strptime(a.period[0], "%Y-%m-%d").date()
        cfg.period_end = datetime.strptime(a.period[1], "%Y-%m-%d").date()
    if a.as_of:
        cfg.as_of = datetime.strptime(a.as_of, "%Y-%m-%d").date()

    out = a.output or a.source.with_name(f"{a.source.stem}_程式產表.xlsx")
    print(f"讀取：{a.source}")
    info = build(a.source, out, cfg)
    s, e = info["period"]
    print(f"報表期間：{s} ～ {e}　基準日：{info['as_of']}")
    print("來源筆數（清理後）：" + "、".join(f"{k} {v}" for k, v in info["rows"].items()))
    n = len(info["log"])
    print(f"資料校正紀錄：{n} 筆" + ("（詳見 _資料校正紀錄）" if n else ""))
    for name, err in info["failed"]:
        print(f"  [警告] {name} 產表失敗：{err}", file=sys.stderr)
    print(f"完成：{out}")


if __name__ == "__main__":
    main()
